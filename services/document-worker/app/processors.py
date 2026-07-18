from __future__ import annotations

import base64
import binascii
import gzip
import hashlib
import html.parser
import io
import json
import os
import re
import subprocess  # noqa: S404 - fixed binary and arguments, no shell
import tarfile
import tempfile
from collections.abc import Callable
from contextlib import suppress
from pathlib import Path, PurePosixPath
from typing import Any, Literal, cast

import fitz  # type: ignore[import-untyped]
import httpx
import pytesseract  # type: ignore[import-untyped]
from defusedxml import ElementTree as safe_element_tree  # type: ignore[import-untyped]
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from .models import CanonicalArtifact, CanonicalBlock, Location, ProcessRequest
from .security import inspect_zip

PROCESSOR_VERSION = "signal-document-worker/1.1.0"
MAX_INPUT_BYTES = 250 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_EXPANDED_BYTES = 500 * 1024 * 1024
MAX_CANONICAL_BLOCKS = 100_000
UNTRUSTED_INSTRUCTION_PATTERNS = (
    re.compile(r"ignore (?:all )?(?:previous|prior|system|developer) instructions", re.I),
    re.compile(
        r"(?:reveal|upload|send|exfiltrate).{0,80}(?:secret|credential|token|password)", re.I | re.S
    ),
)

register_heif_opener()


def _text_block(text: str, **location: Any) -> CanonicalBlock:
    return CanonicalBlock(kind="text", text=text, location=Location(**location))


def _ocr(image: Image.Image) -> tuple[str, dict[str, Any]]:
    normalized = ImageOps.exif_transpose(image).convert("RGB")
    data = pytesseract.image_to_data(normalized, output_type=pytesseract.Output.DICT)
    words: list[str] = []
    confidences: list[float] = []
    for text, confidence in zip(data.get("text", []), data.get("conf", []), strict=False):
        candidate = str(text).strip()
        try:
            score = float(confidence)
        except (TypeError, ValueError):
            score = -1
        if candidate and score >= 0:
            words.append(candidate)
            confidences.append(score)
    text = " ".join(words)
    return text, {
        "ocr_engine": "tesseract",
        "ocr_confidence": round(sum(confidences) / len(confidences), 2) if confidences else None,
        "likely_chart": any(
            marker in text.casefold()
            for marker in ("axis", "legend", "%", "quarter", "revenue", "series")
        ),
    }


def _preview(image: Image.Image) -> str:
    normalized = ImageOps.exif_transpose(image).convert("RGB")
    normalized.thumbnail((1280, 1280))
    output = io.BytesIO()
    normalized.save(output, format="JPEG", quality=82, optimize=True)
    return base64.b64encode(output.getvalue()).decode("ascii")


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    return str(value)


def _parse_pdf(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    reader = PdfReader(io.BytesIO(data), strict=True)
    blocks: list[CanonicalBlock] = []
    warnings: list[str] = []
    preview: str | None = None
    rendered = fitz.open(stream=data, filetype="pdf")
    document_metadata = {
        str(key).lstrip("/"): str(value)[:2_000]
        for key, value in (reader.metadata or {}).items()
        if value is not None
    }
    for index, page in enumerate(reader.pages, start=1):
        visual_page = rendered[index - 1]
        text = page.extract_text() or ""
        links = sorted(
            {
                str(link.get("uri"))
                for link in visual_page.get_links()
                if isinstance(link.get("uri"), str)
            }
        )[:500]
        embedded_images = len(visual_page.get_images(full=True))
        drawing_count = len(visual_page.get_drawings())
        metadata: dict[str, Any] = {
            "extraction": "embedded_text",
            "links": links,
            "embedded_image_count": embedded_images,
            "vector_drawing_count": drawing_count,
            **({"document_metadata": document_metadata} if index == 1 else {}),
        }
        rendered_image: Image.Image | None = None
        if not text.strip():
            pixmap = visual_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            rendered_image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            text, ocr_metadata = _ocr(rendered_image)
            metadata = {**metadata, "extraction": "ocr", **ocr_metadata}
            warnings.append(f"Page {index} required OCR because it contained no embedded text.")
        blocks.append(
            CanonicalBlock(
                kind="text",
                text=text,
                location=Location(page=index),
                metadata=metadata,
            )
        )
        visual_layout_matters = embedded_images > 0 or drawing_count >= 8
        if visual_layout_matters:
            if rendered_image is None:
                pixmap = visual_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                rendered_image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            visual_text, visual_metadata = _ocr(rendered_image)
            blocks.append(
                CanonicalBlock(
                    kind="image",
                    text=visual_text or f"Rendered visual page {index}",
                    location=Location(page=index),
                    metadata={
                        "rendered_page": True,
                        "embedded_image_count": embedded_images,
                        "vector_drawing_count": drawing_count,
                        **visual_metadata,
                    },
                )
            )
            if preview is None:
                preview = _preview(rendered_image)
        if rendered_image is not None:
            rendered_image.close()
    rendered.close()
    return blocks, warnings, preview


def _parse_docx(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    document = Document(io.BytesIO(data))
    blocks = [
        _text_block(paragraph.text, paragraph=index)
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    for table_index, table in enumerate(document.tables, start=1):
        text = "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
        blocks.append(
            CanonicalBlock(
                kind="table",
                text=text,
                location=Location(),
                metadata={"table_index": table_index},
            )
        )
    preview: str | None = None
    warnings: list[str] = []
    seen_parts: set[str] = set()
    for relationship in document.part.rels.values():
        if relationship.is_external:
            continue
        target = relationship.target_part
        content_type = getattr(target, "content_type", "")
        part_name = str(getattr(target, "partname", ""))
        if not str(content_type).startswith("image/") or part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        try:
            with Image.open(io.BytesIO(target.blob)) as image:
                text, metadata = _ocr(image)
                blocks.append(
                    CanonicalBlock(
                        kind="image",
                        text=text or f"Embedded document image {part_name}",
                        location=Location(),
                        metadata={"part_name": part_name, "content_type": content_type, **metadata},
                    )
                )
                if preview is None:
                    preview = _preview(image)
        except (OSError, ValueError):
            warnings.append(f"Embedded image {part_name} could not be decoded.")
    visual_blocks, rendered_preview, render_warnings = _render_office_pages(data, ".docx", "page")
    blocks.extend(visual_blocks)
    warnings.extend(render_warnings)
    return blocks, warnings, rendered_preview or preview


def _chart_title(chart: Any) -> str | None:
    title = getattr(chart, "title", None)
    if title is None:
        return None
    try:
        paragraphs = title.tx.rich.p
        text = " ".join(
            "".join(str(run.t) for run in paragraph.r if getattr(run, "t", None))
            for paragraph in paragraphs
        ).strip()
        return text or None
    except (AttributeError, TypeError):
        return None


def _series_reference(series: Any, names: tuple[str, ...]) -> str | None:
    for name in names:
        container = getattr(series, name, None)
        if container is None:
            continue
        for reference_name in ("strRef", "numRef"):
            reference = getattr(container, reference_name, None)
            formula = getattr(reference, "f", None)
            if isinstance(formula, str) and formula:
                return formula
    return None


def _range_values(workbook: Any, reference: str | None) -> list[Any]:
    if reference is None:
        return []
    match = re.fullmatch(
        r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", reference
    )
    if match is None:
        return []
    sheet_name = (match.group(1) or match.group(2) or "").replace("''", "'")
    if sheet_name not in workbook.sheetnames:
        return []
    cell_range = match.group(3).replace("$", "")
    raw_boundaries = range_boundaries(cell_range)
    if any(value is None for value in raw_boundaries):
        return []
    min_column, min_row, max_column, max_row = cast(tuple[int, int, int, int], raw_boundaries)
    if (max_column - min_column + 1) * (max_row - min_row + 1) > 2_000:
        return []
    sheet = workbook[sheet_name]
    return [
        _json_value(sheet.cell(row=row, column=column).value)
        for row in range(min_row, max_row + 1)
        for column in range(min_column, max_column + 1)
    ]


def _series_title(series: Any, workbook: Any) -> str | None:
    text = getattr(getattr(series, "tx", None), "v", None)
    if isinstance(text, str) and text:
        return text
    reference = _series_reference(series, ("tx",))
    values = _range_values(workbook, reference)
    return str(values[0]) if values else None


def _chart_anchor(chart: Any) -> str | None:
    anchor = getattr(chart, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    try:
        from openpyxl.utils.cell import get_column_letter

        return f"{get_column_letter(int(marker.col) + 1)}{int(marker.row) + 1}"
    except (AttributeError, TypeError, ValueError):
        return None


def _render_office_pages(
    data: bytes, suffix: str, location_kind: str
) -> tuple[list[CanonicalBlock], str | None, list[str]]:
    blocks: list[CanonicalBlock] = []
    preview: str | None = None
    warnings: list[str] = []
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / f"source{suffix}"
        source.write_bytes(data)
        profile = Path(directory) / "profile"
        profile.mkdir(mode=0o700)
        try:
            subprocess.run(  # noqa: S603 - fixed converter, isolated worker, no shell
                [
                    "/usr/bin/libreoffice",
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nolockcheck",
                    "--norestore",
                    "--safe-mode",
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    directory,
                    str(source),
                ],
                check=True,
                capture_output=True,
                timeout=180,
                env={
                    "PATH": "/usr/local/bin:/usr/bin:/bin",
                    "LANG": "C.UTF-8",
                    "SAL_USE_VCLPLUGIN": "svp",
                    "TMPDIR": directory,
                },
            )
            pdf_path = source.with_suffix(".pdf")
            if not pdf_path.is_file() or pdf_path.stat().st_size > MAX_INPUT_BYTES:
                raise ValueError("Office renderer did not produce a bounded PDF.")
            document = fitz.open(pdf_path)
            for index, page in enumerate(document, start=1):
                if index > 200:
                    warnings.append("Office visual rendering was limited to the first 200 pages.")
                    break
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                with Image.open(io.BytesIO(pixmap.tobytes("jpeg"))) as image:
                    text, metadata = _ocr(image)
                    location = (
                        Location(slide=index) if location_kind == "slide" else Location(page=index)
                    )
                    blocks.append(
                        CanonicalBlock(
                            kind="image",
                            text=text or f"Rendered {location_kind} {index}",
                            location=location,
                            metadata={"rendered_office_page": True, **metadata},
                        )
                    )
                    if preview is None:
                        preview = _preview(image)
            document.close()
        except (FileNotFoundError, OSError, subprocess.SubprocessError, ValueError) as error:
            warnings.append(f"Office visual rendering unavailable: {type(error).__name__}.")
    return blocks, preview, warnings


def _convert_office(
    data: bytes,
    source_suffix: str,
    target_suffix: str,
    target_processor: Callable[[bytes], tuple[list[CanonicalBlock], list[str], str | None]],
) -> tuple[list[CanonicalBlock], list[str], str | None]:
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / f"source{source_suffix}"
        source.write_bytes(data)
        profile = Path(directory) / "profile"
        profile.mkdir(mode=0o700)
        target_format = target_suffix.lstrip(".")
        try:
            subprocess.run(  # noqa: S603 - fixed converter, isolated worker, no shell
                [
                    "/usr/bin/libreoffice",
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nolockcheck",
                    "--norestore",
                    "--safe-mode",
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--convert-to",
                    target_format,
                    "--outdir",
                    directory,
                    str(source),
                ],
                check=True,
                capture_output=True,
                timeout=180,
                env={
                    "PATH": "/usr/local/bin:/usr/bin:/bin",
                    "LANG": "C.UTF-8",
                    "SAL_USE_VCLPLUGIN": "svp",
                    "TMPDIR": directory,
                },
            )
        except (FileNotFoundError, OSError, subprocess.SubprocessError) as error:
            raise ValueError(
                f"Safe conversion for {source_suffix} is unavailable: {type(error).__name__}."
            ) from error
        converted = source.with_suffix(target_suffix)
        if not converted.is_file() or converted.stat().st_size > MAX_INPUT_BYTES:
            raise ValueError(f"Safe conversion for {source_suffix} did not produce a bounded file.")
        blocks, warnings, preview = target_processor(converted.read_bytes())
        warnings.insert(
            0,
            f"The original {source_suffix} file was normalized to {target_suffix} "
            "in the isolated worker.",
        )
        return blocks, warnings, preview


def _parse_xlsx(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
    values = load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_links=False)
    blocks: list[CanonicalBlock] = []
    warnings: list[str] = []
    for sheet in formulas.worksheets:
        value_sheet = values[sheet.title]
        blocks.append(
            CanonicalBlock(
                kind="table",
                text=f"Worksheet {sheet.title}: {sheet.calculate_dimension()}",
                location=Location(sheet=sheet.title, range=sheet.calculate_dimension()),
                metadata={
                    "sheet_state": sheet.sheet_state,
                    "merged_ranges": [str(cell_range) for cell_range in sheet.merged_cells.ranges],
                    "tables": [
                        {"name": table.name, "range": table.ref} for table in sheet.tables.values()
                    ],
                    "auto_filter_range": sheet.auto_filter.ref,
                },
            )
        )
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value_cell = value_sheet[cell.coordinate]
                kind: Literal["formula", "text"] = "formula" if cell.data_type == "f" else "text"
                blocks.append(
                    CanonicalBlock(
                        kind=kind,
                        text=str(cell.value),
                        location=Location(sheet=sheet.title, range=cell.coordinate),
                        metadata={"value": value_cell.value, "data_type": cell.data_type},
                    )
                )
                if cell.data_type == "f" and value_cell.value is None:
                    warnings.append(
                        f"{sheet.title}!{cell.coordinate} has no cached formula value; "
                        "recalculate in an office engine."
                    )
        for chart_index, chart in enumerate(getattr(sheet, "_charts", []), start=1):
            series: list[dict[str, Any]] = []
            for raw_series in getattr(chart, "ser", []):
                category_reference = _series_reference(raw_series, ("cat", "xVal"))
                value_reference = _series_reference(raw_series, ("val", "yVal", "bubbleSize"))
                series.append(
                    {
                        "name": _series_title(raw_series, values),
                        "category_reference": category_reference,
                        "categories": _range_values(values, category_reference),
                        "value_reference": value_reference,
                        "values": _range_values(values, value_reference),
                    }
                )
            chart_data = {
                "chart_index": chart_index,
                "chart_type": type(chart).__name__,
                "title": _chart_title(chart),
                "anchor": _chart_anchor(chart),
                "series": series,
            }
            blocks.append(
                CanonicalBlock(
                    kind="table",
                    text=json.dumps(chart_data, ensure_ascii=False, sort_keys=True),
                    location=Location(sheet=sheet.title, range=_chart_anchor(chart)),
                    metadata={"visualization": "chart", **chart_data},
                )
            )
    formulas.close()
    values.close()
    return blocks, sorted(set(warnings)), None


def _parse_pptx(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    presentation = Presentation(io.BytesIO(data))
    blocks: list[CanonicalBlock] = []
    warnings: list[str] = []
    preview: str | None = None
    for slide_index, slide in enumerate(presentation.slides, start=1):
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(
                    CanonicalBlock(
                        kind="text",
                        text=notes,
                        location=Location(slide=slide_index),
                        metadata={"source": "speaker_notes"},
                    )
                )
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str) and text.strip():
                blocks.append(_text_block(text, slide=slide_index))
            if getattr(shape, "has_table", False):
                table_text = "\n".join(
                    "\t".join(cell.text for cell in row.cells) for row in shape.table.rows
                )
                blocks.append(
                    CanonicalBlock(
                        kind="table",
                        text=table_text,
                        location=Location(slide=slide_index),
                        metadata={"shape_name": shape.name},
                    )
                )
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                categories: list[str] = []
                with suppress(AttributeError, IndexError, TypeError):
                    categories = [str(category) for category in chart.plots[0].categories]
                series = [
                    {
                        "name": str(item.name),
                        "values": [_json_value(value) for value in item.values],
                    }
                    for item in chart.series
                ]
                chart_data = {
                    "chart_type": str(chart.chart_type),
                    "categories": categories,
                    "series": series,
                }
                blocks.append(
                    CanonicalBlock(
                        kind="table",
                        text=json.dumps(chart_data, ensure_ascii=False, sort_keys=True),
                        location=Location(slide=slide_index),
                        metadata={"visualization": "chart", "shape_name": shape.name, **chart_data},
                    )
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_blob = shape.image.blob
                    image_checksum = f"sha256:{hashlib.sha256(image_blob).hexdigest()}"
                    with Image.open(io.BytesIO(image_blob)) as image:
                        image_text, image_metadata = _ocr(image)
                        blocks.append(
                            CanonicalBlock(
                                kind="image",
                                text=image_text or f"Slide image {shape.name}",
                                location=Location(slide=slide_index),
                                metadata={
                                    "shape_name": shape.name,
                                    "content_type": shape.image.content_type,
                                    "checksum": image_checksum,
                                    **image_metadata,
                                },
                            )
                        )
                        if preview is None:
                            preview = _preview(image)
                except (OSError, ValueError):
                    warnings.append(f"Slide {slide_index} image {shape.name} could not be decoded.")
    visual_blocks, rendered_preview, render_warnings = _render_office_pages(data, ".pptx", "slide")
    blocks.extend(visual_blocks)
    warnings.extend(render_warnings)
    return blocks, warnings, rendered_preview or preview


def _parse_image(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    with Image.open(io.BytesIO(data)) as image:
        image.verify()
    with Image.open(io.BytesIO(data)) as image:
        normalized = ImageOps.exif_transpose(image)
        ocr_text, ocr_metadata = _ocr(normalized)
        exif = image.getexif()
        block = CanonicalBlock(
            kind="image",
            text=ocr_text or f"Image {normalized.width}×{normalized.height} {normalized.mode}",
            location=Location(),
            metadata={
                "width": normalized.width,
                "height": normalized.height,
                "mode": normalized.mode,
                "orientation_normalized": True,
                "icc_profile_present": "icc_profile" in image.info,
                "exif_present": bool(exif),
                "location_metadata_removed": 34853 in exif,
                **ocr_metadata,
            },
        )
        return [block], [], _preview(normalized)


def _parse_text(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    text = data.decode("utf-8", errors="strict")
    warnings: list[str] = []
    try:
        parsed = json.loads(text)
        normalized = json.dumps(parsed, sort_keys=True, ensure_ascii=False)
        return [_text_block(normalized)], warnings, None
    except json.JSONDecodeError:
        return [_text_block(text)], warnings, None


class _VisibleHtmlParser(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.text: list[str] = []
        self.links: list[str] = []
        self._suppressed = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() in {"script", "style", "template"}:
            self._suppressed += 1
        if tag.casefold() == "a":
            for name, value in attrs:
                if name.casefold() == "href" and value:
                    self.links.append(value[:2_000])

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"script", "style", "template"} and self._suppressed > 0:
            self._suppressed -= 1

    def handle_data(self, data: str) -> None:
        if self._suppressed == 0 and data.strip():
            self.text.append(data.strip())


def _parse_html(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    parser = _VisibleHtmlParser()
    parser.feed(data.decode("utf-8", errors="strict"))
    parser.close()
    return (
        [
            CanonicalBlock(
                kind="text",
                text="\n".join(parser.text),
                location=Location(),
                metadata={"links": sorted(set(parser.links))[:1_000], "scripts_executed": False},
            )
        ],
        [],
        None,
    )


def _parse_xml(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    try:
        root = safe_element_tree.fromstring(data)
    except Exception as error:
        raise ValueError("XML structure is invalid or unsafe.") from error
    blocks: list[CanonicalBlock] = []
    for index, element in enumerate(root.iter(), start=1):
        if index > 100_000:
            raise ValueError("XML contains too many elements.")
        text = (element.text or "").strip()
        if text or element.attrib:
            blocks.append(
                CanonicalBlock(
                    kind="text",
                    text=text or f"<{element.tag}>",
                    location=Location(),
                    metadata={
                        "element": str(element.tag)[:500],
                        "attributes": {
                            str(key)[:500]: str(value)[:2_000]
                            for key, value in element.attrib.items()
                        },
                    },
                )
            )
    return blocks, [], None


def _parse_notebook(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    notebook = json.loads(data.decode("utf-8", errors="strict"))
    if not isinstance(notebook, dict) or not isinstance(notebook.get("cells"), list):
        raise ValueError("Notebook structure is invalid.")
    blocks: list[CanonicalBlock] = []
    for index, cell in enumerate(notebook["cells"], start=1):
        if not isinstance(cell, dict):
            continue
        source = cell.get("source", [])
        text = "".join(source) if isinstance(source, list) else str(source)
        blocks.append(
            CanonicalBlock(
                kind="text",
                text=text,
                location=Location(paragraph=index),
                metadata={
                    "cell_type": str(cell.get("cell_type", "unknown")),
                    "execution_count": _json_value(cell.get("execution_count")),
                    "executed_by_worker": False,
                },
            )
        )
    return blocks, ["Notebook cells were parsed as data and were not executed."], None


def _parse_zip(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    entries = inspect_zip(data)
    blocks = [
        CanonicalBlock(
            kind="archive_entry",
            text=entry.filename,
            location=Location(),
            metadata={"size_bytes": entry.file_size, "compressed_bytes": entry.compress_size},
        )
        for entry in entries
        if not entry.is_dir()
    ]
    return blocks, ["Archive entries were inventoried but uploaded code was not executed."], None


def _parse_tar(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    try:
        with tarfile.open(fileobj=io.BytesIO(data), mode="r:*") as archive:
            members = archive.getmembers()
    except tarfile.TarError as error:
        raise ValueError("TAR archive is invalid.") from error
    if len(members) > MAX_ARCHIVE_ENTRIES:
        raise ValueError("Archive contains too many entries.")
    expanded = sum(max(0, member.size) for member in members)
    if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
        raise ValueError("Archive expansion limit exceeded.")
    blocks: list[CanonicalBlock] = []
    for member in members:
        normalized = PurePosixPath(member.name.replace("\\", "/"))
        if normalized.is_absolute() or ".." in normalized.parts:
            raise ValueError("Archive path traversal was blocked.")
        if member.issym() or member.islnk() or member.isdev():
            raise ValueError("Archive links and device entries are not permitted.")
        if member.isfile():
            blocks.append(
                CanonicalBlock(
                    kind="archive_entry",
                    text=member.name,
                    location=Location(),
                    metadata={"size_bytes": member.size, "mode": member.mode},
                )
            )
    return blocks, ["Archive entries were inventoried but uploaded code was not executed."], None


def _parse_gzip(data: bytes) -> tuple[list[CanonicalBlock], list[str], str | None]:
    try:
        with gzip.GzipFile(fileobj=io.BytesIO(data)) as archive:
            expanded = archive.read(MAX_ARCHIVE_EXPANDED_BYTES + 1)
    except (OSError, EOFError) as error:
        raise ValueError("GZIP archive is invalid.") from error
    if len(expanded) > MAX_ARCHIVE_EXPANDED_BYTES:
        raise ValueError("Archive expansion limit exceeded.")
    try:
        text = expanded.decode("utf-8", errors="strict")
        return [_text_block(text)], ["GZIP content was decompressed as non-executable text."], None
    except UnicodeDecodeError:
        return (
            [
                CanonicalBlock(
                    kind="archive_entry",
                    text="gzip payload",
                    location=Location(),
                    metadata={"size_bytes": len(expanded)},
                )
            ],
            ["Binary GZIP content was inventoried but not executed."],
            None,
        )


def _transcribe(path: str) -> tuple[list[CanonicalBlock], list[str]]:
    endpoint = os.environ.get("TRANSCRIPTION_API_URL")
    if not endpoint:
        return [], ["Media metadata parsed; transcription is not available in this worker."]
    token_file = os.environ.get("TRANSCRIPTION_API_TOKEN_FILE")
    token = ""
    if token_file:
        try:
            token = Path(token_file).read_text(encoding="utf-8").strip()
        except OSError as error:
            return [], [f"Transcription credential unavailable: {type(error).__name__}."]
    elif os.environ.get("NODE_ENV") not in {"production", "staging"}:
        token = os.environ.get("TRANSCRIPTION_API_TOKEN", "")
    headers = {"authorization": f"Bearer {token}"} if token else {}
    model = os.environ.get("TRANSCRIPTION_MODEL", "whisper-1")
    try:
        with Path(path).open("rb") as audio, httpx.Client(timeout=180.0) as client:
            response = client.post(
                endpoint,
                headers=headers,
                data={"model": model, "response_format": "verbose_json"},
                files={"file": (Path(path).name, audio, "application/octet-stream")},
            )
        response.raise_for_status()
        payload = response.json()
        segments = payload.get("segments") if isinstance(payload, dict) else None
        blocks: list[CanonicalBlock] = []
        if isinstance(segments, list):
            for segment in segments:
                if not isinstance(segment, dict) or not isinstance(segment.get("text"), str):
                    continue
                blocks.append(
                    CanonicalBlock(
                        kind="text",
                        text=segment["text"].strip(),
                        location=Location(timestamp_seconds=float(segment.get("start", 0.0))),
                        metadata={
                            "end_seconds": float(segment.get("end", 0.0)),
                            "source": model,
                            **(
                                {"speaker": str(segment["speaker"])}
                                if segment.get("speaker") is not None
                                else {}
                            ),
                        },
                    )
                )
        if not blocks and isinstance(payload, dict) and isinstance(payload.get("text"), str):
            blocks.append(
                CanonicalBlock(
                    kind="text",
                    text=payload["text"].strip(),
                    location=Location(timestamp_seconds=0.0),
                    metadata={"source": model},
                )
            )
        return blocks, [] if blocks else ["The transcription provider returned no transcript."]
    except (OSError, httpx.HTTPError, ValueError, json.JSONDecodeError) as error:
        return [], [f"Transcription unavailable: {type(error).__name__}."]


def _video_frames(path: str, duration: float) -> tuple[list[CanonicalBlock], str | None, list[str]]:
    blocks: list[CanonicalBlock] = []
    warnings: list[str] = []
    preview: str | None = None
    with tempfile.TemporaryDirectory() as directory:
        pattern = str(Path(directory) / "frame-%03d.jpg")
        interval = max(1.0, duration / 8.0)
        try:
            subprocess.run(  # noqa: S603 - fixed executable and bounded output
                [
                    "/usr/bin/ffmpeg",
                    "-v",
                    "error",
                    "-i",
                    path,
                    "-vf",
                    f"fps=1/{interval:.3f}",
                    "-frames:v",
                    "8",
                    pattern,
                ],
                check=True,
                capture_output=True,
                timeout=120,
            )
            for index, frame in enumerate(sorted(Path(directory).glob("frame-*.jpg"))):
                frame_bytes = frame.read_bytes()
                if preview is None:
                    preview = base64.b64encode(frame_bytes).decode("ascii")
                with Image.open(io.BytesIO(frame_bytes)) as image:
                    text, metadata = _ocr(image)
                timestamp = min(duration, index * interval)
                blocks.append(
                    CanonicalBlock(
                        kind="image",
                        text=text or f"Sampled video frame at {timestamp:.3f} seconds",
                        location=Location(timestamp_seconds=timestamp),
                        metadata={"frame_index": index, **metadata},
                    )
                )
        except (FileNotFoundError, subprocess.SubprocessError, OSError) as error:
            warnings.append(f"Video frame analysis unavailable: {type(error).__name__}.")
    return blocks, preview, warnings


def _parse_media(data: bytes, suffix: str) -> tuple[list[CanonicalBlock], list[str], str | None]:
    with tempfile.NamedTemporaryFile(suffix=suffix) as media:
        media.write(data)
        media.flush()
        try:
            result = subprocess.run(  # noqa: S603 - fixed executable, isolated worker image
                [
                    "/usr/bin/ffprobe",
                    "-v",
                    "error",
                    "-show_entries",
                    "format=duration,format_name:stream=codec_type,codec_name,width,height",
                    "-of",
                    "json",
                    media.name,
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=30,
            )
            metadata = json.loads(result.stdout)
            duration = float(metadata.get("format", {}).get("duration", 0.0))
            block = CanonicalBlock(
                kind="media",
                text=f"Media duration {duration:.3f} seconds",
                location=Location(timestamp_seconds=0.0),
                metadata=metadata,
            )
            blocks = [block]
            warnings: list[str] = []
            preview: str | None = None
            streams = metadata.get("streams", [])
            is_video = any(
                isinstance(stream, dict) and stream.get("codec_type") == "video"
                for stream in streams
            )
            transcription_path = media.name
            if is_video:
                frame_blocks, preview, frame_warnings = _video_frames(media.name, duration)
                blocks.extend(frame_blocks)
                warnings.extend(frame_warnings)
                audio_path = f"{media.name}.wav"
                try:
                    subprocess.run(  # noqa: S603 - fixed executable and output path
                        [
                            "/usr/bin/ffmpeg",
                            "-v",
                            "error",
                            "-i",
                            media.name,
                            "-vn",
                            "-ac",
                            "1",
                            "-ar",
                            "16000",
                            "-y",
                            audio_path,
                        ],
                        check=True,
                        capture_output=True,
                        timeout=120,
                    )
                    transcription_path = audio_path
                except (FileNotFoundError, subprocess.SubprocessError):
                    warnings.append("Video audio extraction was unavailable.")
            transcript_blocks, transcript_warnings = _transcribe(transcription_path)
            blocks.extend(transcript_blocks)
            warnings.extend(transcript_warnings)
            if transcription_path != media.name:
                Path(transcription_path).unlink(missing_ok=True)
            return blocks, warnings, preview
        except (FileNotFoundError, subprocess.SubprocessError, json.JSONDecodeError) as error:
            return [], [f"Media inspection unavailable: {type(error).__name__}."], None


Processor = Callable[[bytes], tuple[list[CanonicalBlock], list[str], str | None]]


def process_artifact(request: ProcessRequest) -> CanonicalArtifact:
    try:
        data = base64.b64decode(request.bytes_base64, validate=True)
    except (binascii.Error, ValueError) as error:
        raise ValueError("Artifact bytes are not valid base64.") from error
    if len(data) > MAX_INPUT_BYTES:
        raise ValueError("Artifact exceeds the document worker input limit.")
    actual = f"sha256:{hashlib.sha256(data).hexdigest()}"
    if actual != request.checksum:
        raise ValueError("Artifact checksum mismatch.")

    lower_name = request.file_name.casefold()
    suffix = (
        ".tar.gz"
        if lower_name.endswith(".tar.gz")
        else ".tgz"
        if lower_name.endswith(".tgz")
        else Path(request.file_name).suffix.lower()
    )
    processors: dict[str, Processor] = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".doc": lambda content: _convert_office(content, ".doc", ".docx", _parse_docx),
        ".odt": lambda content: _convert_office(content, ".odt", ".docx", _parse_docx),
        ".rtf": lambda content: _convert_office(content, ".rtf", ".docx", _parse_docx),
        ".xlsx": _parse_xlsx,
        ".xls": lambda content: _convert_office(content, ".xls", ".xlsx", _parse_xlsx),
        ".ods": lambda content: _convert_office(content, ".ods", ".xlsx", _parse_xlsx),
        ".pptx": _parse_pptx,
        ".ppt": lambda content: _convert_office(content, ".ppt", ".pptx", _parse_pptx),
        ".odp": lambda content: _convert_office(content, ".odp", ".pptx", _parse_pptx),
        ".png": _parse_image,
        ".jpg": _parse_image,
        ".jpeg": _parse_image,
        ".webp": _parse_image,
        ".heic": _parse_image,
        ".heif": _parse_image,
        ".gif": _parse_image,
        ".tiff": _parse_image,
        ".bmp": _parse_image,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".json": _parse_text,
        ".jsonl": _parse_text,
        ".csv": _parse_text,
        ".tsv": _parse_text,
        ".yaml": _parse_text,
        ".yml": _parse_text,
        ".html": _parse_html,
        ".htm": _parse_html,
        ".xml": _parse_xml,
        ".ipynb": _parse_notebook,
        ".py": _parse_text,
        ".ts": _parse_text,
        ".js": _parse_text,
        ".sql": _parse_text,
        ".zip": _parse_zip,
        ".tar": _parse_tar,
        ".tar.gz": _parse_tar,
        ".tgz": _parse_tar,
        ".gz": _parse_gzip,
    }
    processor = processors.get(suffix)
    if processor is not None:
        blocks, warnings, preview = processor(data)
    elif suffix in {
        ".mp3",
        ".wav",
        ".m4a",
        ".aac",
        ".flac",
        ".ogg",
        ".mp4",
        ".mov",
        ".webm",
        ".mkv",
    }:
        blocks, warnings, preview = _parse_media(data, suffix)
    else:
        raise ValueError(f"Unsupported document format: {suffix or request.mime_type}")
    if len(blocks) > MAX_CANONICAL_BLOCKS:
        raise ValueError("Artifact produced too many canonical blocks.")
    if any(
        pattern.search(block.text) for block in blocks for pattern in UNTRUSTED_INSTRUCTION_PATTERNS
    ):
        warnings.append(
            "Untrusted content contains instruction-like text; treat it only as document data."
        )
    return CanonicalArtifact(
        artifact_id=request.artifact_id,
        original_checksum=request.checksum,
        detected_format=suffix.lstrip(".") or request.mime_type,
        processor_version=PROCESSOR_VERSION,
        blocks=blocks,
        warnings=warnings,
        preview_base64=preview,
    )
