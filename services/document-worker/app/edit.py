from __future__ import annotations

import base64
import hashlib
import io
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from .models import (
    AddSlideOperation,
    AppendParagraphOperation,
    AppendRowOperation,
    EditRequest,
    GeneratedArtifact,
    ReplaceTextOperation,
    SetCellOperation,
)


def _replace(value: str, operation: ReplaceTextOperation) -> tuple[str, int]:
    if operation.replace_all:
        return value.replace(operation.find, operation.replacement), value.count(operation.find)
    if operation.find not in value:
        return value, 0
    return value.replace(operation.find, operation.replacement, 1), 1


def _edit_docx(data: bytes, request: EditRequest) -> tuple[bytes, dict[str, Any], list[str]]:
    document = Document(io.BytesIO(data))
    replacements = 0
    for operation in request.operations:
        if isinstance(operation, ReplaceTextOperation):
            for paragraph in document.paragraphs:
                paragraph.text, count = _replace(paragraph.text, operation)
                replacements += count
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell.text, count = _replace(cell.text, operation)
                        replacements += count
        elif isinstance(operation, AppendParagraphOperation):
            document.add_paragraph(operation.text)
        else:
            raise ValueError(f"Operation {operation.type} is not valid for a Word document.")
    output = io.BytesIO()
    document.save(output)
    checked = Document(io.BytesIO(output.getvalue()))
    return (
        output.getvalue(),
        {
            "paragraphs": len(checked.paragraphs),
            "tables": len(checked.tables),
            "replacements": replacements,
        },
        [],
    )


def _edit_xlsx(data: bytes, request: EditRequest) -> tuple[bytes, dict[str, Any], list[str]]:
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    changed_cells: list[str] = []
    for operation in request.operations:
        if isinstance(operation, SetCellOperation):
            if operation.sheet not in workbook.sheetnames:
                raise ValueError(f"Worksheet {operation.sheet} does not exist.")
            workbook[operation.sheet][operation.cell] = operation.value
            changed_cells.append(f"{operation.sheet}!{operation.cell}")
        elif isinstance(operation, AppendRowOperation):
            if operation.sheet not in workbook.sheetnames:
                raise ValueError(f"Worksheet {operation.sheet} does not exist.")
            target_sheet = workbook[operation.sheet]
            target_sheet.append(operation.values)
            changed_cells.append(f"{operation.sheet}!{target_sheet.max_row}:{target_sheet.max_row}")
        elif isinstance(operation, ReplaceTextOperation):
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str):
                            updated, count = _replace(cell.value, operation)
                            if count > 0:
                                cell.value = updated
                                changed_cells.append(f"{worksheet.title}!{cell.coordinate}")
        else:
            raise ValueError(f"Operation {operation.type} is not valid for a workbook.")
    output = io.BytesIO()
    workbook.save(output)
    checked = load_workbook(io.BytesIO(output.getvalue()), data_only=False, read_only=True)
    return (
        output.getvalue(),
        {"sheets": checked.sheetnames, "changed_cells": sorted(set(changed_cells))},
        [],
    )


def _edit_pptx(data: bytes, request: EditRequest) -> tuple[bytes, dict[str, Any], list[str]]:
    presentation = Presentation(io.BytesIO(data))
    replacements = 0
    for operation in request.operations:
        if isinstance(operation, ReplaceTextOperation):
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.text, count = _replace(run.text, operation)
                            replacements += count
        elif isinstance(operation, AddSlideOperation):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.6), Inches(0.7))
            title_box.text = operation.title
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(4.8))
            body.text = operation.body
        else:
            raise ValueError(f"Operation {operation.type} is not valid for a presentation.")
    output = io.BytesIO()
    presentation.save(output)
    checked = Presentation(io.BytesIO(output.getvalue()))
    return output.getvalue(), {"slides": len(checked.slides), "replacements": replacements}, []


def _edit_text(data: bytes, request: EditRequest) -> tuple[bytes, dict[str, Any], list[str]]:
    text = data.decode("utf-8", errors="strict")
    replacements = 0
    for operation in request.operations:
        if isinstance(operation, ReplaceTextOperation):
            text, count = _replace(text, operation)
            replacements += count
        elif isinstance(operation, AppendParagraphOperation):
            text = f"{text.rstrip()}\n\n{operation.text}\n"
        else:
            raise ValueError(f"Operation {operation.type} is not valid for a text artifact.")
    return text.encode(), {"characters": len(text), "replacements": replacements}, []


def edit_artifact(request: EditRequest) -> GeneratedArtifact:
    data = base64.b64decode(request.bytes_base64, validate=True)
    actual = f"sha256:{hashlib.sha256(data).hexdigest()}"
    if actual != request.checksum:
        raise ValueError("Artifact checksum mismatch.")
    suffix = Path(request.file_name).suffix.lower()
    if suffix == ".docx":
        output, preview, warnings = _edit_docx(data, request)
    elif suffix == ".xlsx":
        output, preview, warnings = _edit_xlsx(data, request)
    elif suffix == ".pptx":
        output, preview, warnings = _edit_pptx(data, request)
    elif suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml"}:
        output, preview, warnings = _edit_text(data, request)
    else:
        raise ValueError(f"Versioned editing is not supported for {suffix or request.mime_type}.")
    checksum = f"sha256:{hashlib.sha256(output).hexdigest()}"
    mime_types = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    return GeneratedArtifact(
        file_name=request.output_file_name,
        mime_type=mime_types.get(suffix, request.mime_type),
        checksum=checksum,
        bytes_base64=base64.b64encode(output).decode("ascii"),
        preview=preview,
        checks=["source_checksum_verified", "reopened_generated_file", "checksum_recorded"],
        warnings=warnings,
    )
